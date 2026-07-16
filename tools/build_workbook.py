#!/usr/bin/env python3
"""Build a single .xlsx (one tab per category) for upload to Google Sheets.

Drag the file into Google Drive -> it opens as a Sheet with every tab intact.
When the schema settles, the same workbook imports straight into Airtable.

Data is EVALUATED out of the app via JavaScriptCore, not regex-scraped, so the
workbook is exactly what the app loads.

Usage: python3 tools/build_workbook.py [path-to-html] [out.xlsx]
"""
import re, sys, json, subprocess, os
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter

JSC = "/System/Library/Frameworks/JavaScriptCore.framework/Versions/A/Helpers/jsc"
SRC = sys.argv[1] if len(sys.argv) > 1 else "/Users/andy/Downloads/BFT_Living_Database_6.html"
OUT = sys.argv[2] if len(sys.argv) > 2 else os.path.join(os.path.dirname(__file__), "..", "export", "BFT_Living_Database.xlsx")
SKIP = {"network_person", "network_company", "network_connection"}  # backend-only, per instruction

NAVY, GOLD = "1F3864", "BF9000"
ARIAL = "Arial"

def load():
    s = re.search(r"<script>\n(.*?)\n</script>", open(SRC, encoding="utf-8").read(), re.S).group(1)
    a = s.find("const SEED_DATA"); b = s.find("\n  };", a) + 4
    open("/tmp/_wb.js", "w", encoding="utf-8").write(
        s[a:b].replace("const SEED_DATA", "var SEED_DATA", 1) + "\nprint(JSON.stringify(SEED_DATA));")
    r = subprocess.run([JSC, "/tmp/_wb.js"], capture_output=True, text=True)
    if not r.stdout.strip():
        raise SystemExit("failed to evaluate SEED_DATA: " + r.stderr[:300])
    return json.loads(r.stdout)

def sheet_from(wb, name, rows):
    ws = wb.create_sheet(name[:31])
    cols = []
    for r in rows:
        for k in r:
            if k not in cols: cols.append(k)
    # bftId and name lead -- they are the identity of the row
    for lead in ("name", "bftId"):
        if lead in cols: cols.insert(0, cols.pop(cols.index(lead)))
    ws.append(cols)
    for c in range(1, len(cols) + 1):
        cell = ws.cell(row=1, column=c)
        cell.font = Font(name=ARIAL, bold=True, color="FFFFFF", size=10)
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    for r in rows:
        ws.append([r.get(c, "") for c in cols])
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name=ARIAL, size=10)
            cell.alignment = Alignment(vertical="top", wrap_text=False)
    for idx, c in enumerate(cols, 1):
        longest = max([len(str(c))] + [len(str(r.get(c, ""))) for r in rows[:200]])
        ws.column_dimensions[get_column_letter(idx)].width = min(max(longest + 2, 10), 55)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{len(rows) + 1}"
    ws.row_dimensions[1].height = 28
    return ws

def readme(wb, data):
    ws = wb.create_sheet("READ ME", 0)
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 104
    rows = [
        ("BFT Living Intelligence Database", ""),
        ("", ""),
        ("What this is", "A snapshot of the live database at https://bft-living-database.vercel.app, one tab per category."),
        ("Generated", "16 Jul 2026, by tools/build_workbook.py"),
        ("", ""),
        ("HOW TO USE IN GOOGLE SHEETS", ""),
        ("Upload", "Drag this .xlsx into Google Drive. It opens as a Sheet with every tab intact. No conversion needed."),
        ("Editing", "Safe to sort, filter, and edit. Sorting cannot scramble identity because every row carries a bftId."),
        ("", ""),
        ("THE bftId COLUMN -- DO NOT EDIT OR DELETE", ""),
        ("Why it exists", "Rows previously had no permanent identity -- they were keyed by NAME. Names collide (Mike Garcia is not Chuy Garcia) and change (David Ellison / Edison). Any sync keyed on a name silently merges or loses people."),
        ("What it does", "bftId is permanent and unique. It is the join key back to the app, and later into Airtable. Delete it and a row can no longer be matched to its source."),
        ("Adding rows", "Leave bftId BLANK on new rows. It gets assigned on import -- do not invent one."),
        ("", ""),
        ("MOVING TO AIRTABLE LATER", ""),
        ("How", "File > Download > CSV per tab, then import each into Airtable. Or import this .xlsx directly. bftId carries across as the join key."),
        ("Record limits", "Airtable free = 1,000 records per base; this snapshot is already ~713. Team ($20/user/mo) = 50,000. Google Sheets has no comparable ceiling -- which is why Sheets suits the build phase."),
        ("", ""),
        ("SOURCING RULES (see the criteria tab)", ""),
        ("Hierarchy", "Primary filings/official records > the subject's own publication > named reporting > Wikipedia (navigation only, NEVER the citation). Wikipedia is not used anywhere in this database."),
        ("sourceUrl", "A clickable, verified URL. Blank means UNSOURCED -- treat the claim as unverified, not as true."),
        ("", ""),
        ("NOT INCLUDED", ""),
        ("LinkedIn network", "The network_person (5,834) / network_company (1,299) tables are deliberately excluded -- backend-only, per standing instruction."),
        ("Contact information", "Not present. If contact details get added, this file becomes PII about named private individuals alongside stance ratings -- treat sharing accordingly."),
        ("", ""),
        ("CONTENTS", ""),
    ]
    for cat, r in data.items():
        if cat not in SKIP: rows.append((cat, f"{len(r)} records"))
    for a, b in rows: ws.append([a, b])
    for i, (a, b) in enumerate(rows, 1):
        ac, bc = ws.cell(row=i, column=1), ws.cell(row=i, column=2)
        head = a and not b
        ac.font = Font(name=ARIAL, bold=bool(head or i == 1), size=13 if i == 1 else 10,
                       color=GOLD if head and i > 1 else (NAVY if i == 1 else "000000"))
        bc.font = Font(name=ARIAL, size=10)
        bc.alignment = Alignment(wrap_text=True, vertical="top")
        ac.alignment = Alignment(vertical="top")
    return ws

def main():
    data = load()
    wb = Workbook(); wb.remove(wb.active)
    n = 0
    for cat, rows in data.items():
        if cat in SKIP or not rows: continue
        sheet_from(wb, cat, rows); n += len(rows)
        print(f"  {cat:<14s} {len(rows):>4d} rows")
    readme(wb, data)
    os.makedirs(os.path.dirname(os.path.abspath(OUT)), exist_ok=True)
    wb.save(OUT)
    print(f"\n{n} records across {len(wb.sheetnames)-1} tabs -> {os.path.abspath(OUT)}")

if __name__ == "__main__":
    main()
